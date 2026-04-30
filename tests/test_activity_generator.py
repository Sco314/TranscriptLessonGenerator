"""Offline tests for web/activity_generator.

Validates the JSON schema and renderer without making any network calls.
A fixture spec mirroring tests/smoke_test.py is rendered and inspected.
"""

from __future__ import annotations

import sys
from pathlib import Path

import jsonschema
import pytest
from docx import Document

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "web"))

import activity_generator as ag  # noqa: E402


FIXTURE_SPEC = {
    "theme": "teal",
    "header_name": "Smoke Test Worksheet",
    "page1": [
        {"type": "title", "text": "Customer Service Excellence"},
        {"type": "subtitle", "text": "Indeed Career Tips Series — Smoke Test"},
        {"type": "video_link", "url": "https://youtu.be/EXAMPLE"},
        {"type": "name_date_line"},
        {
            "type": "info_box",
            "title": "The STAR Framework",
            "paragraphs": [
                [
                    {"text": "Situation  ", "bold": True, "color": "C62828"},
                    {"text": "— set the scene"},
                ],
                [
                    {"text": "Task  ", "bold": True, "color": "EF6C00"},
                    {"text": "— what you needed to do"},
                ],
            ],
        },
        {"type": "section_header", "text": "Watch For These Strategies"},
        {
            "type": "para",
            "text": "Fill in the blanks as you watch the video.",
            "italic": True,
            "color": "helper_text",
        },
        {
            "type": "numbered_question",
            "number": 1,
            "timestamp": "0:42",
            "runs": [
                {"text": "The video says to start by listening for the "},
                {"text": "real problem", "bold": True, "color": "C62828"},
                {"text": ", not just the "},
                {"blank": 18},
                {"text": "."},
            ],
        },
        {
            "type": "scaffolded_prompt",
            "label": "In your own words, what does 'service recovery' mean?",
            "hint": "(1–2 sentences)",
        },
        {
            "type": "tip_box",
            "text": "Empathy doesn't mean agreement — it means you heard them.",
        },
    ],
    "page2": [
        {"type": "section_header", "text": "Score the Candidates"},
        {
            "type": "candidates_list",
            "candidates": ["Trevor", "Kimberly", "Wayne"],
            "role": "Customer Service Rep",
        },
        {
            "type": "evaluation_table",
            "candidates": ["Trevor", "Kimberly", "Wayne"],
            "criteria": ["Listened?", "Acknowledged?", "Solution?"],
            "questions": [
                {"text": "Tell me about a time you handled an angry customer.", "timestamp": "2:30"},
                {"text": "What does great service mean to you?", "timestamp": "4:15"},
            ],
        },
        {
            "type": "question_header",
            "text": "Reflection: Which candidate would you hire?",
            "timestamp": "end",
        },
        {
            "type": "warning_box",
            "text": "Be honest in your scoring — this isn't about who you'd be friends with.",
        },
        {"type": "section_header", "text": "Apply It Yourself"},
        {"type": "job_selection_line", "default_job": "Customer Service Rep"},
        {
            "type": "numbered_question",
            "number": 5,
            "runs": [{"text": "Describe a time YOU handled a frustrated person."}],
        },
        {
            "type": "self_check_box",
            "items": [
                "I named the strategies (not just numbered them)",
                "I scored every candidate on every criterion",
            ],
        },
    ],
}


def test_schema_accepts_fixture():
    jsonschema.validate(FIXTURE_SPEC, ag.WORKSHEET_SCHEMA)


def test_schema_rejects_unknown_theme():
    bad = {**FIXTURE_SPEC, "theme": "neon"}
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(bad, ag.WORKSHEET_SCHEMA)


def test_schema_rejects_missing_pages():
    bad = {"theme": "teal", "header_name": "x", "page1": [{"type": "title", "text": "x"}]}
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(bad, ag.WORKSHEET_SCHEMA)


def test_schema_rejects_unknown_component():
    bad = {
        "theme": "teal",
        "header_name": "x",
        "page1": [{"type": "fancy_thing", "text": "hi"}],
        "page2": [{"type": "title", "text": "x"}],
    }
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(bad, ag.WORKSHEET_SCHEMA)


def test_render_produces_docx(tmp_path):
    output = tmp_path / "test.docx"
    ok = ag.render(FIXTURE_SPEC, output)
    assert ok, "render() should return True on success"
    assert output.exists() and output.stat().st_size > 5000

    # Sanity-check that python-docx can reopen what we just wrote.
    doc = Document(str(output))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Customer Service Excellence" in text
    assert "Watch For These Strategies" in text
    assert "Apply It Yourself" in text
    # Tables: candidates_list adds paragraphs; evaluation_table + boxes add tables.
    assert len(doc.tables) >= 4, f"expected at least 4 tables, got {len(doc.tables)}"


def test_build_briefing_uses_lesson_fields():
    class _L:
        lesson_id = "ted_test"
        title = "Test Lesson"
        youtube_url = "https://youtu.be/X"
        ted_url = "https://ed.ted.com/lessons/test"
        transcript = "blah"
        author = "Author"
        collection = "Coll"
        duration = "5:00"

    briefing = ag.build_briefing(_L(), {"format": "interview_game", "theme": "navy"})
    assert briefing["VIDEO"] == "Test Lesson"
    assert briefing["URL"] == "https://youtu.be/X"
    assert briefing["FORMAT"] == "interview_game"
    assert briefing["THEME_COLOR"] == "navy"
    assert briefing["TRANSCRIPT"] == "blah"
