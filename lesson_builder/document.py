"""
document.py — Document factory and save helper.

create_document() sets up the page (US Letter, 0.5" margins, optional running
header with doc name + page numbers) and gives you a Document instance to
hand to the building-block helpers in components.py.

save_document() wraps doc.save() with the project's ✅/❌ console convention
so failures don't fail silently.
"""

from pathlib import Path
from typing import Optional, Union

from docx import Document
from docx.document import Document as _DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from .constants import COLOR, FONT, PAGE, SIZE


def create_document(*, header_name: Optional[str] = None,
                    use_header: bool = False) -> _DocumentType:
    """Create a fresh Document with US Letter, 0.5" margins, Arial 10pt default.

    use_header=True + header_name="..." enables a right-aligned header reading:
        <header_name>  |  Page X of Y
    The top margin is bumped to 0.625" to make room for the header.
    """
    doc = Document()

    # ── Default font (Arial 10pt) ──
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    from docx.shared import Pt
    style.font.size = Pt(SIZE["body"] / 2.0)

    # ── Page setup ──
    section = doc.sections[0]
    section.page_width = Twips(PAGE["width"])
    section.page_height = Twips(PAGE["height"])
    section.left_margin = Twips(PAGE["margins"]["left"])
    section.right_margin = Twips(PAGE["margins"]["right"])
    section.bottom_margin = Twips(PAGE["margins"]["bottom"])
    section.top_margin = Twips(
        PAGE["header_top_margin"] if use_header else PAGE["margins"]["top"]
    )

    if use_header and header_name:
        _add_running_header(section, header_name)

    return doc


def _add_running_header(section, doc_name: str):
    """Right-aligned header: '<doc_name>  |  Page X of Y'."""
    header = section.header
    # Use the first existing paragraph (python-docx adds one by default)
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    from docx.shared import Pt

    label = p.add_run(f"{doc_name}  |  Page ")
    label.font.name = FONT
    label.font.size = Pt(SIZE["page_number"] / 2.0)
    label.font.color.rgb = _color(COLOR["helper_text"])

    # PAGE field
    _add_field(p, "PAGE")

    of = p.add_run(" of ")
    of.font.name = FONT
    of.font.size = Pt(SIZE["page_number"] / 2.0)
    of.font.color.rgb = _color(COLOR["helper_text"])

    # NUMPAGES field
    _add_field(p, "NUMPAGES")


def _add_field(paragraph, instr: str):
    """Insert a Word field (PAGE, NUMPAGES, etc.) into a paragraph."""
    from docx.shared import Pt

    run = paragraph.add_run()
    run.font.name = FONT
    run.font.size = Pt(SIZE["page_number"] / 2.0)
    run.font.color.rgb = _color(COLOR["helper_text"])
    r = run._element

    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    r.append(fldChar1)

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = f" {instr} "
    r.append(instrText)

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    r.append(fldChar2)


def _color(hex_str: str):
    from docx.shared import RGBColor
    return RGBColor.from_string(hex_str.upper().lstrip("#"))


def save_document(doc: _DocumentType, filepath: Union[str, Path]) -> bool:
    """Save the document. Logs ✅ on success, ❌ with reason on failure."""
    path = Path(filepath)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
        print(f"✅ Document saved: {path}")
        return True
    except PermissionError as e:
        print(f"❌ Could not save document: file is open or locked. ({e})")
        return False
    except OSError as e:
        print(f"❌ Could not save document: filesystem error. ({e})")
        return False
    except Exception as e:
        print(f"❌ Unexpected error saving document: {type(e).__name__}: {e}")
        return False
